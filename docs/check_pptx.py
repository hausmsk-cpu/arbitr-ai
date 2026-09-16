"""Проверка структуры сгенерированной презентации."""

from pathlib import Path

from pptx import Presentation

PPTX = Path(r"C:\Users\МЕГАЛИТ\Desktop\AI Dip\Arbitr AI\docs\Arbitr AI_презентация_защита.pptx")

prs = Presentation(str(PPTX))
print(f"Слайдов: {len(prs.slides)} | Размер: {prs.slide_width.inches:.1f}×{prs.slide_height.inches:.1f} in\n")

for i, slide in enumerate(prs.slides, 1):
    texts = []
    has_chart = False
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            first = sh.text_frame.text.strip().split("\n")[0]
            if not texts and (len(first) > 4):
                texts.append(first[:65])
        if sh.has_chart:
            has_chart = True
    label = texts[0] if texts else "(схема/таблица)"
    flag = " 📊" if has_chart else ""
    print(f"{i:2}. {label}{flag}")