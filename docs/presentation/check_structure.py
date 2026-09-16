"""Проверка структуры презентации защиты Arbitr AI."""

from pathlib import Path

from pptx import Presentation

PPTX = Path(__file__).parent.parent / "Arbitr_AI_презентация_защита.pptx"


def main() -> None:
    prs = Presentation(str(PPTX))
    print(f"Слайдов: {len(prs.slides)}\n")

    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        has_img = False
        has_table = False
        for sh in slide.shapes:
            if sh.shape_type == 13:  # PICTURE
                has_img = True
            if sh.has_table:
                has_table = True
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t.split("\n")[0][:70])
        label = texts[0] if texts else "(без явного текста)"
        badges = []
        if has_img:
            badges.append("🖼")
        if has_table:
            badges.append("📋")
        print(f"{i:2}. {label}  {''.join(badges)}")


if __name__ == "__main__":
    main()