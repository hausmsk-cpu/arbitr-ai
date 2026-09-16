"""Проверка структуры презентации Arbitr AI к защите."""

from pathlib import Path

from pptx import Presentation

PPTX = Path(r"C:\Users\МЕГАЛИТ\Desktop\AI Dip\Arbitr AI\docs\Arbitr AI_презентация_защита.pptx")


def main() -> None:
    prs = Presentation(str(PPTX))
    print(f"Слайдов: {len(prs.slides)}\n")

    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        has_img = False
        has_table = False
        for sh in slide.shapes:
            if sh.shape_type == 13:  # PICTURE
                has_img = True
            if sh.has_table:
                has_table = True
            if sh.has_text_frame and sh.text_frame.text.strip():
                t = sh.text_frame.text.strip().splitlines()[0].strip()
                if t and len(t) > 1:
                    texts.append(t[:70])

        label = texts[0] if texts else "(графика / без подписи)"
        badges = []
        if has_img:
            badges.append("🖼")
        if has_table:
            badges.append("📋")
        print(f"{i:2}. {label}  {''.join(badges)}")


if __name__ == "__main__":
    main()